"""Gera "miniprojeto 1 doc.docx" na raiz (para abrir no Google Docs) e as figuras em relatorio/figuras/.

Uso (na raiz do repositório, depois de extrair_dados.py):
    python tools/relatorio/gerar_docx.py
Requer python-docx (pip install python-docx).
"""

import json
import os
import re

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.ticker as mticker
import numpy as np
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from matplotlib.colors import to_hex, to_rgb
from matplotlib.patches import Rectangle

OUT_DIR = "relatorio"
FIG = os.path.join(OUT_DIR, "figuras")
os.makedirs(FIG, exist_ok=True)
D = json.load(open(os.path.join(OUT_DIR, "dados", "doc_data.json")))
CV = json.load(open(os.path.join(OUT_DIR, "dados", "curves.json")))
M, C, AUG, CH, HY = D["mlp"], D["cnn"], D["aug"], D["champions"], D["hyper"]
PT = {"airplane": "avião", "automobile": "automóvel", "bird": "pássaro", "cat": "gato", "deer": "cervo",
      "dog": "cachorro", "frog": "sapo", "horse": "cavalo", "ship": "navio", "truck": "caminhão"}
VAL = "val/accuracy_mean"

# ============================================================================ estilo dos gráficos
INK, INK2, MUTED, HAIR, SURF, ACCENT = "#121722", "#465063", "#6f7889", "#dce1e9", "#ffffff", "#0d6b66"
MLP_C, CNN_C = "#2a78d6", "#eb6834"
SERIES = ["#4a3aa7", "#1baf7a", "#eda100"]
SEQ = ["#e6eefa", "#cde2fb", "#9ec5f4", "#6da7ec", "#2a78d6", "#1c5cab", "#0d366b"]
plt.rcParams.update({
    "font.family": "DejaVu Sans", "font.size": 9, "axes.edgecolor": HAIR, "axes.labelcolor": INK2,
    "xtick.color": MUTED, "ytick.color": MUTED, "axes.grid": False, "figure.facecolor": SURF,
    "axes.facecolor": SURF, "savefig.facecolor": SURF, "axes.spines.top": False, "axes.spines.right": False,
})


def pct(x, d=1):
    return "—" if x is None else f"{x * 100:.{d}f}%".replace(".", ",")


def sd(x):
    return "" if x is None else f"± {x * 100:.1f}".replace(".", ",")


def pp(x, d=1):
    return ("+" if x >= 0 else "−") + f"{abs(x) * 100:.{d}f}".replace(".", ",") + " p.p."


def num(x, d=1):
    return "—" if x is None else f"{x:.{d}f}".replace(".", ",")


def params(n):
    return f"{n / 1e6:.2f} M".replace(".", ",") if n >= 1e6 else f"{n / 1e3:.0f} mil"


def save(fig, name):
    path = os.path.join(FIG, name)
    fig.savefig(path, dpi=220, bbox_inches="tight")
    plt.close(fig)
    return path


def hgrid(ax, axis="x"):
    ax.grid(True, axis=axis, color=HAIR, linewidth=0.8)
    ax.set_axisbelow(True)


def pct_axis(axis):
    axis.set_major_formatter(mticker.FuncFormatter(lambda t, _: f"{t * 100:.0f}%"))


# ============================================================================ gráficos
def progression(stages, color, name):
    fig, ax = plt.subplots(figsize=(6.4, 0.42 * len(stages) + 0.7))
    ys = np.arange(len(stages))[::-1]
    for y, s in zip(ys, stages):
        v, e = s["v"], s["sd"]
        if s.get("bonus"):
            ax.axhline(y + 0.5, color=HAIR, lw=0.8)
        ax.plot([v - e, v + e], [y, y], color=color, lw=2, solid_capstyle="round")
        ax.scatter([v], [y], s=46, color=SURF if s.get("bonus") else color, edgecolor=color if s.get("bonus") else SURF,
                   linewidth=1.8 if s.get("bonus") else 1.5, zorder=3)
        ax.text(v + e + 0.008, y, pct(v), va="center", fontsize=9, color=INK, fontweight="bold")
    ax.set_yticks(ys, [f"{s['label']}\n{s['sub']}" for s in stages], fontsize=8.5, color=INK2)
    ax.set_xlim(0.35, 0.9)
    pct_axis(ax.xaxis)
    hgrid(ax)
    ax.spines["left"].set_visible(False)
    ax.tick_params(axis="y", length=0)
    ax.set_xlabel("acurácia no teste (média ± desvio em 5 folds) · ponto vazado = bônus", fontsize=8.5)
    return save(fig, name)


def champ_row(net):
    return next(r for r in AUG[net]["rows"] if r["exp"] == AUG[net]["champ"])


mlp_st = [("Linear (B0)", "regressão logística"), ("MLP original (B0)", "64-128-64"), ("Topologia (B1)", "3 × 256"),
          ("Otimização (B2)", "+ Adam lr 3e-4"), ("Checagem", "4 × 256"), ("Regularização (B3)", "+ dropout 0,2")]
cnn_st = [("LeNet original (B0)", "2 blocos, 32/64 filtros"), ("Topologia (B1)", "4 blocos, k3, maxpool"),
          ("Otimização (B2)", "Adam lr 1e-3 (mantido)"), ("Regularização (B3)", "+ BatchNorm, dropout 0,5"),
          ("Pooling (+B)", "3 blocos, pooling 2×2")]
FIG_PROG_MLP = progression([{"label": l, "sub": s, "v": x["test"], "sd": x["test_std"]} for (l, s), x in zip(mlp_st, M["stages"])]
                           + [{"label": "Bônus: augmentation", "sub": "crop + flip, 150 épocas", "v": champ_row("mlp")["test"],
                               "sd": champ_row("mlp")["test_std"], "bonus": True}], MLP_C, "evolucao_mlp.png")
FIG_PROG_CNN = progression([{"label": l, "sub": s, "v": x["test"], "sd": x["test_std"]} for (l, s), x in zip(cnn_st, C["stages"])]
                           + [{"label": "Bônus: augmentation", "sub": "crop + flip, 4 blocos", "v": champ_row("cnn")["test"],
                               "sd": champ_row("cnn")["test_std"], "bonus": True}], CNN_C, "evolucao_cnn.png")


def heatmap_ax(ax, rows, cols, cell, domain, row_fmt, col_fmt, show_rows=True):
    for i, r in enumerate(rows):
        for j, c in enumerate(cols):
            d = cell(r, c)
            if d is None:
                ax.add_patch(Rectangle((j - 0.46, i - 0.44), 0.92, 0.88, facecolor="#f6f8fb", edgecolor=HAIR, hatch="////", linewidth=0))
                ax.text(j, i, "—", ha="center", va="center", color=MUTED, fontsize=9)
                continue
            v, champ = d
            k = min(6, int(min(1, max(0, (v - domain[0]) / (domain[1] - domain[0]))) * 7))
            ax.add_patch(Rectangle((j - 0.46, i - 0.44), 0.92, 0.88, facecolor=SEQ[k],
                                   edgecolor=INK if champ else "none", linewidth=2 if champ else 0))
            ax.text(j, i, pct(v), ha="center", va="center", fontsize=8.5, fontweight="bold", color="#ffffff" if k >= 4 else INK)
    ax.set_xlim(-0.5, len(cols) - 0.5)
    ax.set_ylim(len(rows) - 0.5, -0.5)
    ax.set_xticks(range(len(cols)), [col_fmt(c) for c in cols], fontsize=8)
    ax.set_yticks(range(len(rows)), [row_fmt(r) for r in rows] if show_rows else [""] * len(rows), fontsize=8)
    ax.tick_params(length=0)
    for sp in ax.spines.values():
        sp.set_visible(False)


def scale_note(fig, domain, note=""):
    fig.text(0.01, -0.04, f"Escala de cor: {pct(domain[0], 0)} (claro) → {pct(domain[1], 0)} (escuro). Contorno = campeão. {note}",
             fontsize=7.5, color=MUTED)


def heatmap(name, rows, cols, cell, domain, row_fmt=str, col_fmt=str, size=(6.2, 2.6), note=""):
    fig, ax = plt.subplots(figsize=size)
    heatmap_ax(ax, rows, cols, cell, domain, row_fmt, col_fmt)
    scale_note(fig, domain, note)
    return save(fig, name)


def facets(name, panels, rows, cols, domain, row_fmt, col_fmt, size, note=""):
    fig, axes = plt.subplots(1, len(panels), figsize=size, gridspec_kw={"wspace": 0.08})
    for i, (ax, (title, cell)) in enumerate(zip(axes, panels)):
        heatmap_ax(ax, rows, cols, cell, domain, row_fmt, col_fmt, show_rows=(i == 0))
        ax.set_title(title, fontsize=9, color=INK2, loc="left")
    scale_note(fig, domain, note)
    return save(fig, name)


def find(rows, pred):
    return next((r for r in rows if pred(r)), None)


def lr_fmt(lr):
    return f"{lr:g}".replace(".", ",")


LRS = [0.0001, 0.0003, 0.001, 0.003, 0.01, 0.03, 0.1]
FIG_HM_MLP_B1 = heatmap("heatmap_mlp_b1.png", [1, 2, 3, 4, 5], [128, 256, 512, 1024, 2048],
    lambda r, c: (lambda x: x and (x[VAL], x["label"] == "L3_N256"))(find(M["b1_triagem"], lambda o: o["mlp_layers"] == r and o["mlp_neurons"] == c)),
    (0.495, 0.535), lambda r: f"{r} camada{'s' if r > 1 else ''}", lambda c: f"{c} neurônios", size=(6.4, 3.0))
FIG_HM_MLP_B2 = heatmap("heatmap_mlp_b2.png", ["adam", "sgd"], LRS,
    lambda r, c: (lambda x: x and (x[VAL], r == "adam" and c == 0.0003))(find(M["b2_triagem"], lambda o: o["optimizer"] == r and abs(o["lr"] - c) < 1e-12)),
    (0.42, 0.54), str, lr_fmt, size=(6.6, 1.6), note="Abaixo de 42% = cor mais clara (10% = colapso).")
FIG_HM_MLP_B3 = heatmap("heatmap_mlp_b3.png", [0, 0.2, 0.3, 0.5], ["cross_entropy", "mse"],
    lambda r, c: (lambda x: x and (x[VAL], r == 0.2 and c == "cross_entropy"))(find(M["b3_triagem"], lambda o: abs(o["dropout"] - r) < 1e-9 and o["loss_fn"] == c)),
    (0.525, 0.555), lambda r: f"dropout {r:g}".replace(".", ","), lambda c: "MSE" if c == "mse" else "entropia cruzada", size=(4.2, 2.4))
FIG_HM_CNN_B1 = facets("heatmap_cnn_b1.png",
    [(title, (lambda red: lambda b, kp: (lambda x: x and (x[VAL], x["label"] == "B4_k3_padsame_maxpool"))(
        find(C["b1_triagem"], lambda o: o["conv_blocks"] == b and o["kernel_size"] == kp[0] and o["padding"] == kp[1] and o["reducao"] == red)))(red))
     for red, title in [("maxpool", "Max pooling 2×2"), ("stride2", "Convolução com stride 2")]],
    [2, 3, 4], [(3, "same"), (3, "valid"), (5, "same"), (5, "valid")], (0.65, 0.765),
    lambda b: f"{b} blocos", lambda kp: f"k{kp[0]} · {kp[1]}", (7.4, 2.3), "Hachurado = combinação descartada (mapa 0×0).")
FIG_HM_CNN_B2 = heatmap("heatmap_cnn_b2.png", ["adam", "sgd"], LRS,
    lambda r, c: (lambda x: x and (x[VAL], r == "adam" and c == 0.001))(find(C["b2_triagem"], lambda o: o["optimizer"] == r and abs(o["lr"] - c) < 1e-12)),
    (0.64, 0.76), str, lr_fmt, size=(6.6, 1.6), note="Abaixo de 64% = cor mais clara.")
FIG_HM_CNN_B3 = heatmap("heatmap_cnn_b3.png", [0, 0.3, 0.5], [False, True],
    lambda r, c: (lambda x: x and (x[VAL], r == 0.5 and c))(find(C["b3_triagem"], lambda o: abs(o["cnn_dropout"] - r) < 1e-9 and o["cnn_batch_norm"] == c and o["pool_size"] == 2)),
    (0.745, 0.8), lambda r: f"dropout {r:g}".replace(".", ","), lambda c: "com BatchNorm" if c else "sem BatchNorm", size=(4.2, 2.0))
FIG_HM_CNN_B = heatmap("heatmap_cnn_pooling.png", [2, 3, 4], [2, 3, 4],
    lambda b, p: (lambda x: x and (x[VAL], x["label"] == "B3_pool2"))(find(C["extraB"], lambda o: o["label"] == f"B{b}_pool{p}")),
    (0.745, 0.8), lambda b: f"{b} blocos", lambda p: f"pooling {p}×{p}", size=(4.2, 2.1), note="Hachurado = mapa 0×0.")
FIG_HM_AUG_CNN = facets("heatmap_augmentation_cnn.png",
    [(title, (lambda aug: lambda b, d: (lambda x: x and (x["val"], x["exp"] == AUG["cnn"]["champ"]))(
        find(AUG["cnn"]["rows"], lambda o: o["augment"] == aug and o["conv_blocks"] == b and abs(o["dropout"] - d) < 1e-9)))(aug))
     for aug, title in [(False, "Sem augmentation"), (True, "Com augmentation")]],
    [3, 4], [0.3, 0.5], (0.78, 0.875), lambda b: f"{b} blocos", lambda d: f"dropout {d:g}".replace(".", ","), (6.4, 1.8),
    "Folds: 5 para campeão e referência; 3 nas demais.")
FIG_HM_AUG_MLP = heatmap("heatmap_augmentation_mlp.png", [0, 0.2], [False, True],
    lambda d, a: (lambda x: x and (x["val"], x["exp"] == AUG["mlp"]["champ"]))(find(AUG["mlp"]["rows"], lambda o: o["augment"] == a and abs(o["dropout"] - d) < 1e-9)),
    (0.53, 0.58), lambda d: f"dropout {d:g}".replace(".", ","), lambda a: "com augmentation" if a else "sem augmentation", size=(4.2, 1.6))


def curves(key, name, show_train=True):
    spec = CV[key]
    fig, axes = plt.subplots(1, 2, figsize=(7.2, 2.7))
    idx = {"accuracy": (1, 3), "loss": (5, 7)}
    for ax, metric in zip(axes, ["accuracy", "loss"]):
        ti, vi = idx[metric]
        for i, sr in enumerate(spec["series"]):
            pts = np.array(sr["points"])
            if show_train:
                ax.plot(pts[:, 0], pts[:, ti], color=SERIES[i], lw=1.6, ls=(0, (4, 3)), alpha=0.9)
            ax.plot(pts[:, 0], pts[:, vi], color=SERIES[i], lw=2, label=sr["label"])
            best = pts[np.argmax(pts[:, vi])] if metric == "accuracy" else pts[np.argmin(pts[:, vi])]
            ax.scatter([best[0]], [best[vi]], s=26, color=SERIES[i], edgecolor=SURF, linewidth=1.2, zorder=3)
        for sr in spec.get("context", []):
            pts = np.array(sr["points"])
            ax.plot(pts[:, 0], pts[:, vi], color=MUTED, lw=2, label=sr["label"])
        hgrid(ax, "y")
        ax.set_xlabel("época", fontsize=8.5)
        ax.set_title("Acurácia" if metric == "accuracy" else "Loss", fontsize=9, color=INK2, loc="left")
        ax.xaxis.set_major_locator(mticker.MaxNLocator(integer=True))
        if metric == "accuracy":
            pct_axis(ax.yaxis)
        else:
            ax.yaxis.set_major_formatter(mticker.FuncFormatter(lambda t, _: f"{t:.1f}".replace(".", ",")))
    handles, labels = axes[0].get_legend_handles_labels()
    if show_train:
        handles += [plt.Line2D([], [], color=MUTED, lw=2), plt.Line2D([], [], color=MUTED, lw=1.6, ls=(0, (4, 3)))]
        labels += ["validação", "treino"]
    ncol = 3 if len(labels) > 5 else min(5, len(labels))
    fig.legend(handles, labels, loc="upper center", ncol=ncol, frameon=False, fontsize=8, bbox_to_anchor=(0.5, 1.12 if ncol == 3 else 1.08))
    fig.tight_layout()
    return save(fig, name)


FIG_CV = {k: curves(k, f"curvas_{k}.png", show_train=(k != "mlp_b2")) for k in CV}


def dumbbell(name, left, right, left_label, right_label, left_color, right_color, xlim=(0.3, 1.0)):
    cls = sorted([(c, left[i], right[i]) for i, c in enumerate(D["classes"])], key=lambda t: t[2] - t[1], reverse=True)
    fig, ax = plt.subplots(figsize=(6.6, 3.6))
    ys = np.arange(len(cls))[::-1]
    for y, (c, a, b) in zip(ys, cls):
        ax.plot([a, b], [y, y], color="#c5ccd7", lw=2, zorder=1)
        ax.scatter([a], [y], s=42, color=left_color, edgecolor=SURF, linewidth=1.4, zorder=3)
        ax.scatter([b], [y], s=42, color=right_color, edgecolor=SURF, linewidth=1.4, zorder=3)
        ax.text(1.005, y, pp(b - a, 0), va="center", fontsize=8.5, color=INK, fontweight="bold", transform=ax.get_yaxis_transform())
    ax.set_yticks(ys, [PT[c] for c, _, _ in cls], fontsize=8.5, color=INK2)
    ax.set_xlim(*xlim)
    pct_axis(ax.xaxis)
    hgrid(ax)
    ax.spines["left"].set_visible(False)
    ax.tick_params(axis="y", length=0)
    ax.set_xlabel("recall no teste (média de 5 folds)", fontsize=8.5)
    ax.legend(handles=[plt.Line2D([], [], marker="o", ls="", color=left_color, label=left_label),
                       plt.Line2D([], [], marker="o", ls="", color=right_color, label=right_label)],
              loc="lower center", bbox_to_anchor=(0.45, 1.0), ncol=2, frameon=False, fontsize=8.5)
    return save(fig, name)


FIG_CLASSES = dumbbell("recall_por_classe.png", M["recall_final"], C["recall_final"], "MLP 4×256 (55,2%)", "CNN 3 blocos (79,5%)", MLP_C, CNN_C)

def mix(color, t):
    """Mistura `color` com branco: t = 0 → branco, t = 1 → cor pura."""
    a, b = np.array(to_rgb(color)), np.array(to_rgb(SURF))
    return to_hex(b + (a - b) * t)


def f1_stages(name):
    cls = D["classes"] + ["macro"]
    fig, axes = plt.subplots(1, 2, figsize=(8.8, 4.3), gridspec_kw={"wspace": 0.06, "width_ratios": [7, 6]})
    dom = (0.3, 0.95)
    for i, (ax, net, title) in enumerate(zip(axes, ["mlp", "cnn"], ["MLP", "CNN"])):
        cols = [c for c in CH if c["net"] == net]
        heatmap_ax(ax, cls, cols,
                   lambda r, c: ((c["overall"]["test"]["f1_macro"]["mean"] if r == "macro" else c["classes"]["test"]["f1"]["mean"][D["classes"].index(r)]),
                                 "campeã" in c["label"]),
                   dom, lambda r: "média macro" if r == "macro" else PT[r], lambda c: c["short"], show_rows=(i == 0))
        ax.set_title(title, fontsize=9, color=INK2, loc="left")
        for t in ax.texts:
            t.set_fontsize(8)
            t.set_text(t.get_text().replace("%", ""))
    scale_note(fig, dom, "Valores = F1 no teste em %, média de 5 folds.")
    return save(fig, name)


def confusion_fig(ch, name):
    cm = np.array(ch["confusion"])
    rows = cm.sum(axis=1)
    fig, ax = plt.subplots(figsize=(5.4, 4.9))
    for i in range(10):
        for j in range(10):
            f = cm[i, j] / rows[i]
            if i == j:
                t = 0.18 + 0.72 * min(1, max(0, (f - 0.2) / 0.75))
                face, ink, weight = mix(ACCENT, t), ("#ffffff" if t > 0.52 else INK), "bold"
            elif f >= 0.005:
                k = min(6, int(min(1, f / 0.2) * 7))
                face, ink, weight = SEQ[k], ("#ffffff" if k >= 4 else INK), "normal"
            else:
                face, ink, weight = "#f6f8fb", MUTED, "normal"
            ax.add_patch(Rectangle((j - 0.47, i - 0.47), 0.94, 0.94, facecolor=face, linewidth=0))
            ax.text(j, i, f"{round(f * 100)}", ha="center", va="center", fontsize=7.5, color=ink, fontweight=weight)
    ax.set_xlim(-0.5, 9.5); ax.set_ylim(9.5, -0.5)
    ax.set_xticks(range(10), [PT[c] for c in D["classes"]], rotation=45, ha="right", fontsize=8)
    ax.set_yticks(range(10), [PT[c] for c in D["classes"]], fontsize=8)
    ax.set_xlabel("classe prevista", fontsize=8.5); ax.set_ylabel("classe real", fontsize=8.5)
    ax.tick_params(length=0)
    for sp in ax.spines.values():
        sp.set_visible(False)
    return save(fig, name)


FIG_F1_STAGES = f1_stages("f1_por_classe_etapas.png")
FIG_CM = {ch["exp"]: confusion_fig(ch, f"confusao_{ch['exp']}.png") for ch in CH}

# ============================================================================ documento
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.0)
USABLE = Cm(17)
BODY_FONT, HEAD_FONT = "IBM Plex Sans", "IBM Plex Sans Condensed"
st = doc.styles["Normal"]
st.font.name, st.font.size = BODY_FONT, Pt(10.5)
st.font.color.rgb = RGBColor.from_string("121722")
st.paragraph_format.space_after, st.paragraph_format.line_spacing = Pt(6), 1.15
for name, size, color in (("Title", 26, "121722"), ("Heading 1", 18, "121722"), ("Heading 2", 14, "121722"), ("Heading 3", 11.5, "0d6b66")):
    hs = doc.styles[name]
    hs.font.name, hs.font.size, hs.font.bold = HEAD_FONT, Pt(size), True
    hs.font.color.rgb = RGBColor.from_string(color)
    rpr = hs.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), HEAD_FONT)
    hs.paragraph_format.space_before = Pt({"Title": 0, "Heading 1": 18, "Heading 2": 14, "Heading 3": 10}[name])
    hs.paragraph_format.space_after, hs.paragraph_format.keep_with_next = Pt(6), True
fig_counter = [0]


def rich(p, text, size=None, color=None, italic=False):
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "IBM Plex Mono"; r.font.size = Pt(9)
        else:
            r = p.add_run(part)
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        if italic:
            r.italic = True


def para(text="", size=None, color=None, italic=False, style=None):
    p = doc.add_paragraph(style=style)
    if text:
        rich(p, text, size, color, italic)
    return p


def question(text):
    return para(text, italic=True, color="465063")


def bullets(items, size=None):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        rich(p, it, size)
        p.paragraph_format.space_after = Pt(3)


def figure(path, caption, width=USABLE):
    fig_counter[0] += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=width)
    cap = doc.add_paragraph()
    rich(cap, f"Figura {fig_counter[0]} — {caption}", size=8.5, color="6f7889", italic=True)
    cap.paragraph_format.space_after = Pt(10)


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def table(header, rows, widths=None, highlight=None, font=9):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(font - 0.5); r.font.color.rgb = RGBColor.from_string("465063")
        shade(c, "EEF1F5")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            rich(cells[i].paragraphs[0], str(v), size=font)
            cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)
            if highlight and highlight(ri):
                shade(cells[i], "DFF0EE")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def callout(label, text):
    t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"
    c = t.rows[0].cells[0]; shade(c, "F6F8FB"); c.text = ""
    p = c.paragraphs[0]
    r = p.add_run(label.upper() + "  "); r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor.from_string("0d6b66")
    rich(p, text, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def facts(items):
    table([k for k, _ in items], [[v for _, v in items]], font=9.5)


def slides(items):
    p = para("Para os slides", size=9, color="0d6b66"); p.runs[0].bold = True
    p.paragraph_format.space_after = Pt(2); p.paragraph_format.keep_with_next = True
    for it in items:
        q = doc.add_paragraph(style="List Bullet"); rich(q, it, size=9); q.paragraph_format.space_after = Pt(2)


def note(text):
    para(text, size=9.5, color="465063")


def stage(s):
    return [pct(s["val"]) + " " + sd(s["val_std"]), pct(s["test"]) + " " + sd(s["test_std"]), params(s["params"])]


def rank_row(r, name):
    return [name, pct(r[VAL]) + " " + sd(r["val/accuracy_std"]), num(r["gap/accuracy_mean"], 3), num(r["melhor_epoca_media"]), params(r["num_parameters"])]


def class_table(ch):
    """Precision, recall e F1 por classe em treino, validação e teste (média de 5 folds; ± do F1 no teste)."""
    splits = [("train", "Treino"), ("val", "Validação"), ("test", "Teste")]
    t = doc.add_table(rows=2, cols=11)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    def head(cell, text):
        cell.text = ""
        r = cell.paragraphs[0].add_run(text); r.bold = True; r.font.size = Pt(7.5); r.font.color.rgb = RGBColor.from_string("465063")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        shade(cell, "EEF1F5")

    top = t.rows[0].cells
    top[0].merge(t.rows[1].cells[0]); head(top[0], "Classe")
    for k, (_, name) in enumerate(splits):
        a = top[1 + 3 * k].merge(top[3 + 3 * k]); head(a, name)
        for m, lab in enumerate(("P", "R", "F1")):
            head(t.rows[1].cells[1 + 3 * k + m], lab)
    top[10].merge(t.rows[1].cells[10]); head(top[10], "± F1 teste")
    S = ch["classes"]
    f1t = S["test"]["f1"]["mean"]
    worst = int(np.argmin(f1t))
    for i, c in enumerate(D["classes"]):
        cells = t.add_row().cells
        vals = [PT[c]] + [num(S[sp][m]["mean"][i] * 100) for sp, _ in splits for m in ("precision", "recall", "f1")] + [num(S["test"]["f1"]["std"][i] * 100)]
        for j, v in enumerate(vals):
            cells[j].text = ""
            r = cells[j].paragraphs[0].add_run(v); r.font.size = Pt(8)
            if j == 9:
                r.bold = True
            cells[j].paragraphs[0].paragraph_format.space_after = Pt(0)
            if j:
                cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if i == worst:
                shade(cells[j], "FBEFD6")
    O = ch["overall"]
    cells = t.add_row().cells
    vals = ["Macro"] + [num(O[sp][k]["mean"] * 100) for sp, _ in splits for k in ("precision_macro", "recall_macro", "f1_macro")] + [num(O["test"]["f1_macro"]["std"] * 100)]
    for j, v in enumerate(vals):
        cells[j].text = ""
        r = cells[j].paragraphs[0].add_run(v); r.font.size = Pt(8); r.bold = True
        cells[j].paragraphs[0].paragraph_format.space_after = Pt(0)
        if j:
            cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        shade(cells[j], "EEF1F5")
    for row in t.rows:
        for j, w in enumerate([2.3] + [1.4] * 9 + [1.6]):
            row.cells[j].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ---------------------------------------------------------------- capa e resumo
p = para("UFPE · REDES NEURAIS · ESTUDO DE ABLAÇÃO NO CIFAR-10", size=9, color="0d6b66"); p.runs[0].bold = True
doc.add_paragraph("Caderno de Resultados CIFAR-10", style="Title")
para("Todos os resultados de MLP e CNN, bloco a bloco: o que cada experimento mostrou, que gráfico usar e o que dizer nos slides. "
     "Documento preliminar para o grupo montar a apresentação.", size=11.5, color="465063")
para("432 runs no Weights & Biases · blocos 0–3, complementares e bônus (data augmentation) concluídos.", size=9, color="6f7889")

doc.add_heading("Resumo", level=1)
table(["", "Melhor MLP", "Melhor CNN"], [
    ["Acurácia no teste", "**55,2% ± 0,3**", "**79,5% ± 0,4**"],
    ["Configuração", "4 camadas × 256 neurônios, Adam lr 3e-4, dropout 0,2", "3 blocos convolucionais, pooling 2×2, BatchNorm, dropout 0,5"],
    ["Ganho sobre a referência", "+4,0 p.p. sobre a MLP original (51,1%)", "+9,8 p.p. sobre a LeNet original (69,7%)"],
    ["Com data augmentation (bônus)", "**57,8% ± 0,2** (+2,7 p.p. pareado)", "**86,4% ± 0,6** (+7,4 p.p. pareado, 4 blocos)"],
], widths=[4, 6.5, 6.5], font=9.5)
figure(FIG_PROG_MLP, "Evolução da MLP por etapa: acurácia no teste (média ± desvio em 5 folds). Ponto vazado = bônus.", width=Cm(14.5))
figure(FIG_PROG_CNN, "Evolução da CNN por etapa: acurácia no teste (média ± desvio em 5 folds). Ponto vazado = bônus.", width=Cm(14.5))
callout("Tese", "A MLP para em ~55% porque recebe a imagem achatada e perde a vizinhança entre pixels: decide muito por cor e fundo. "
        "A CNN aproveita a estrutura espacial por construção e ganha em todas as classes (+24,3 p.p. no teste). Dentro de cada rede, "
        "os maiores ganhos vieram de **regularização** (dropout na MLP, BatchNorm na CNN); a otimização rendeu pouco ou nada. "
        "Como bônus, **data augmentation** acrescentou **+7,4 p.p.** à CNN e só **+2,7 p.p.** à MLP.")

# ---------------------------------------------------------------- protocolo
doc.add_heading("Como os números foram obtidos", level=1)
para("O mesmo protocolo vale para todos os blocos. Vale um slide só para ele: é o que dá credibilidade às comparações.")
table(["Etapa", "Descrição"], [
    ["Dados", "CIFAR-10: 50.000 imagens de treino e 10.000 de teste, 32×32×3, 10 classes balanceadas. O arquivo usado no Kaggle foi conferido por MD5 contra o oficial."],
    ["Validação cruzada", "5 folds estratificados sobre o treino (40.000 / 10.000 por fold), **com as mesmas partições em todos os experimentos** (seed 42). As comparações podem ser feitas fold a fold."],
    ["Grid search em blocos", "Cada bloco varia um grupo de hiperparâmetros e herda automaticamente o campeão do bloco anterior. Nenhum valor foi escolhido à mão."],
    ["Triagem e confirmação", "Toda configuração roda os folds 1–3; as 3 melhores completam os folds 4–5. O campeão é a maior acurácia de validação média nos 5 folds."],
    ["Seleção × teste", "Toda escolha usa só a **validação**. O teste é revelado apenas para os campeões, no fim (evita vazamento de dados)."],
    ["Treino", "Early stopping na `val/loss`, restauração dos melhores pesos, métricas de treino em modo avaliação para o gap treino–validação ser justo com dropout."],
    ["Curvas por época", "Média entre folds, só nas épocas em que todos os folds ainda treinavam: a curva termina quando o primeiro fold aciona o early stopping. Mesmos valores registrados no W&B (passo = época)."],
    ["Registro", "Cada época vai para `historico_treino.csv` (loss, acurácia, precision, recall e F1 gerais e por classe) antes de ser espelhada no W&B."],
    ["Infraestrutura", "Kaggle, 2× NVIDIA T4, dados inteiros na GPU. Cada notebook final executa o estudo completo da rede: ~1h–1h30 (MLP) e ~3h–3h30 (CNN)."],
], widths=[3.6, 13.4], font=9)

# ---------------------------------------------------------------- MLP
doc.add_heading("MLP · Perceptron multicamadas", level=1)
para("Imagem achatada em 3.072 entradas. Ordem dos blocos: referências → topologia → otimização → checagem → regularização e função de erro → complementar → bônus.")

doc.add_heading("Bloco 0 · Referências", level=2)
question("De onde partimos? Mesmo protocolo dos blocos, aplicado a um classificador linear e à MLP do notebook original.")
table(["Configuração", "Validação", "Teste", "Parâmetros"], [["Regressão logística (0 camadas)", *stage(M["stages"][0])], ["MLP original `64-128-64`", *stage(M["stages"][1])]],
      widths=[6.5, 3.5, 3.5, 3.5], highlight=lambda i: i == 1)
bullets(["O classificador linear fica em **38,8%** no teste: é o piso sem nenhuma não linearidade.",
         "A MLP original chega a **51,1%**. Esse é o número contra o qual todo ganho da MLP é medido."])

doc.add_heading("Bloco 1 · Topologia", level=2)
question("Qual combinação de profundidade e largura absorve melhor o CIFAR-10? Adam 1e-3, sem regularização, early stopping com paciência 5. "
         "Grid: camadas {1, 2, 3, 4, 5} × neurônios {128, 256, 512, 1024, 2048} = 25 configurações.")
figure(FIG_HM_MLP_B1, "MLP, Bloco 1: acurácia de validação na triagem (média dos folds 1–3).", width=Cm(14))
figure(FIG_CV["mlp_b1"], "MLP, Bloco 1: curvas por época (média dos folds 1–3). Redes maiores abrem mais cedo a distância entre treino e validação.")
bullets(["**Profundidade ajuda até certo ponto:** sair de 1 para 2+ camadas rende +1,5 a 2,5 p.p.",
         "**Largura acima de ~256 não ajuda.** Redes largas atingem o ótimo mais cedo (época 3–4) e passam a decorar o treino.",
         "Quase todas as configurações de 2 a 5 camadas ficam entre 51,5% e 53,1%: diferenças do tamanho do desvio. A frase honesta é “profundidade moderada basta”.",
         "**Maldição do vencedor ao vivo:** o 1º da triagem tinha 53,1% em 3 folds e caiu para 52,7% em 5."])
facts([("Campeão", "3 × 256"), ("Validação · 5 folds", "52,7% ± 0,7"), ("Teste", "52,5% ± 0,3"), ("Finalistas empatados", "4×256 · 2×256")])
slides(["Figura principal: `outputs_mlp/_relatorio/heatmap_mlp_b1_topologia_mlp_layers_x_mlp_neurons__val-accuracy_mean.png`",
        "W&B: painel de linha, x = `epoch`, y = `val/accuracy`, agrupar por Group, grupos `mlp_b1_topologia__L1_N256`, `__L3_N256`, `__L5_N2048`."])

doc.add_heading("Bloco 2 · Otimização", level=2)
question("Com a topologia 3×256 fixa, qual algoritmo e taxa de aprendizagem convergem melhor? Grid: {SGD, Adam} × lr {1e-4 … 1e-1} = 14 configurações. SGD usa momentum 0,9.")
figure(FIG_HM_MLP_B2, "MLP, Bloco 2: acurácia de validação na triagem.")
figure(FIG_CV["mlp_b2"], "MLP, Bloco 2: validação por época (média dos folds 1–3). Em cinza, o colapso do Adam com lr 3e-2.")
bullets(["**Adam** funciona de 1e-4 a 1e-3 e **colapsa em 10%** (acaso) a partir de lr 3e-2: a rede passa a prever uma única classe.",
         "**SGD** só funciona entre 1e-3 e 3e-2, mas nessa faixa empata com Adam (lr 1e-2 = 53,3%).",
         "Adam com lr 3e-4 venceu: +0,8 p.p. no teste sobre o Bloco 1."])
facts([("Campeão", "Adam · lr 3e-4"), ("Validação · 5 folds", "53,4% ± 0,5"), ("Teste", "53,3% ± 0,3"), ("Convergência", "Adam: época 5–7")])
note("**Ressalvas.** SGD com lr ≤ 3e-4 usou as 40 épocas inteiras e ainda estava aprendendo: o resultado ruim é limitado pelo orçamento. "
     "O colapso em 10% não aparece como “divergência” (a loss não virou NaN); nos slides, chame de colapso.")
slides(["Figura principal: `outputs_mlp/_relatorio/heatmap_mlp_b2_otimizacao_optimizer_x_lr__val-accuracy_mean.png`", "Velocidade: `…__melhor_epoca_media.png`."])

doc.add_heading("Checagem de interação", level=2)
question("A melhor topologia com Adam 1e-3 continua melhor com o otimizador novo? O 2º e o 3º do Bloco 1 foram treinados com Adam 3e-4 nos 5 folds.")
table(["Configuração · 5 folds", "Validação", "Gap", "Melhor época", "Parâmetros"],
      [rank_row(M["b2_final"][0], "Campeão do Bloco 2 · `3 × 256`")] +
      [rank_row(r, f"{'2º' if r['label'].startswith('rank2') else '3º'} do Bloco 1 · `{r['mlp_layers']} × 256`") for r in M["chk_final"]],
      widths=[6, 3, 2.5, 2.5, 3], highlight=lambda i: i == 1)
bullets(["4×256 superou 3×256 por apenas 0,1 p.p., dentro do ruído, e virou a base do Bloco 3.",
         "As topologias do topo são **estatisticamente equivalentes**: a busca em blocos não deixou passar nenhuma rede claramente melhor."])

doc.add_heading("Bloco 3 · Regularização e função de erro", level=2)
question("Com a rede 4×256 convergindo bem, como a taxa de dropout e a função de erro afetam a generalização? Grid: dropout {0; 0,2; 0,3; 0,5} × {entropia cruzada, MSE}. 50 épocas, paciência 7.")
figure(FIG_HM_MLP_B3, "MLP, Bloco 3: acurácia de validação na triagem.", width=Cm(10))
figure(FIG_CV["mlp_b3"], "MLP, Bloco 3: curvas por época (média dos folds 1–3). O dropout empurra a melhor época para depois.")
bullets(["**Dropout 0,2 deu o maior ganho de uma única etapa na MLP: +1,5 p.p.**, bem acima do desvio.",
         "O gap treino–validação **aumentou** (0,144 contra 0,134): o dropout não eliminou o overfitting, ele o **atrasou** (melhor época 15 contra 7), permitindo mais épocas de aprendizado útil.",
         "MSE ficou ~0,3 p.p. abaixo da entropia cruzada e converge mais devagar (época 20 contra 15).",
         "Dropout 0,5 causa subajuste (ver Complementar A)."])
facts([("Campeão da MLP", "dropout 0,2 · entropia cruzada"), ("Validação · 5 folds", "55,0% ± 0,6"), ("Teste", "55,2% ± 0,3"), ("Ganho sobre o original", "+4,0 p.p.")])
slides(["Figura principal: `outputs_mlp/_relatorio/heatmap_mlp_b3_regularizacao_dropout_x_loss_fn__val-accuracy_mean.png`",
        "Argumento do “dropout atrasa o overfitting”: as curvas por época acima."])

doc.add_heading("Complementar A · Dropout com mais épocas", level=2)
question("No Bloco 3, dropout 0,5 atingiu a melhor época perto do limite de 50. Com 100 épocas e paciência 10, ele alcança o 0,2?")
pa = {p["exp"]: p for p in M["paired_A"]}
table(["Configuração · 5 folds", "Validação", "Melhor época", "vs campeão (por fold)"],
      [[f"dropout {r['dropout']:g} · {'MSE' if r['loss_fn'] == 'mse' else 'entropia cruzada'}".replace(".", ","), pct(r[VAL]) + " " + sd(r["val/accuracy_std"]),
        num(r["melhor_epoca_media"]), f"{pp(pa[r['exp_name']]['delta'], 2)} · venceu {pa[r['exp_name']]['wins']}/5" if r["exp_name"] in pa else "referência"]
       for r in sorted(M["extraA"], key=lambda r: -r[VAL])],
      widths=[6, 3.5, 3, 4.5], highlight=lambda i: i == 0)
bullets(["**O limite de épocas cortou pouco:** com o dobro de épocas, dropout 0,5 melhorou só +0,3 p.p.",
         "Dropout 0,5 perdeu para 0,2 em **5 de 5 folds** (−1,25 p.p.): subajuste real, não falta de tempo.",
         "Entropia cruzada venceu MSE em **5 de 5 folds** (+0,34 p.p.): pequeno na média, consistente fold a fold.",
         "O campeão da MLP não muda."])

# ---------------------------------------------------------------- CNN
doc.add_heading("CNN · Rede convolucional", level=1)
para("Cada bloco convolucional é `Conv2d → (BatchNorm2d) → ReLU → MaxPool`, com 64 filtros no primeiro bloco dobrando a cada bloco, seguido de uma camada densa de 512.")

doc.add_heading("Bloco 0 · Referência", level=2)
question("LeNet-5 adaptada do notebook original: 2 blocos (32 e 64 filtros), kernel 3×3, camadas densas 120 e 84.")
table(["Configuração", "Validação", "Teste", "Parâmetros"], [["LeNet adaptada", *stage(C["stages"][0])]], widths=[6.5, 3.5, 3.5, 3.5])
bullets(["Mesmo pequena (0,5 M parâmetros), a LeNet chega a **69,7%** no teste: +14,5 p.p. sobre a melhor MLP, que tem o dobro de parâmetros."])

doc.add_heading("Bloco 1 · Topologia e extração espacial", level=2)
question("Quantos blocos, qual janela de convolução, qual padding e como reduzir a resolução: max pooling 2×2 ou convolução com stride 2? "
         "Grid: blocos {2, 3, 4} × kernel {3, 5} × padding {same, valid} × redução {maxpool, stride 2} = 24 combinações, 20 válidas.")
figure(FIG_HM_CNN_B1, "CNN, Bloco 1: acurácia de validação na triagem, com a mesma escala de cor nos dois painéis.")
figure(FIG_CV["cnn_b1"], "CNN, Bloco 1: curvas por época (média dos folds 1–3).")
bullets(["**O resultado mais nítido dos blocos:** max pooling venceu stride 2 em **todas** as combinações. A pior configuração com pooling (70,6%) supera a melhor com stride (68,5%).",
         "Mais profundidade ajudou (2 → 3 → 4 blocos com kernel 3 e padding same).",
         "Kernel 3×3 venceu 5×5 mesmo com menos parâmetros. Padding same ficou levemente à frente de valid.",
         "4 combinações foram descartadas antes do treino: padding valid com redes profundas ou kernel 5 reduz o mapa a 0×0."])
facts([("Campeão", "4 blocos · k3 · same · maxpool"), ("Validação · 5 folds", "76,1% ± 0,8"), ("Teste", "75,5% ± 0,6"), ("Ganho sobre a LeNet", "+5,8 p.p.")])
slides(["Figura principal: `outputs_cnn/_relatorio/heatmap_cnn_b1_topologia_conv_blocks_x_kernel_size-padding__val-accuracy_mean.png`",
        "Custo: `…__num_parameters.png` mostra que kernel 5 custa mais e rende menos."])

doc.add_heading("Bloco 2 · Otimização", level=2)
question("Com a topologia de 4 blocos fixa, qual algoritmo e taxa de aprendizagem convergem melhor? Mesmo grid da MLP.")
figure(FIG_HM_CNN_B2, "CNN, Bloco 2: acurácia de validação na triagem.")
bullets(["**Não houve ganho:** o Adam padrão (lr 1e-3) já era adequado; Adam 3e-4 e SGD 1e-2 empataram. É um achado legítimo.",
         "Adam com lr 3e-3 já fica instável; em 1e-2 ou mais colapsa para ~10–22%.",
         "SGD com lr ≤ 1e-3 foi cortado pelo limite de 40 épocas."])
facts([("Campeão", "Adam · lr 1e-3"), ("Validação · 5 folds", "75,8% ± 0,4"), ("Teste", "75,7% ± 0,2"), ("Ganho sobre o Bloco 1", "+0,3 p.p. (ruído)")])
slides(["`outputs_cnn/_relatorio/heatmap_cnn_b2_otimizacao_optimizer_x_lr__val-accuracy_mean.png`, lado a lado com o da MLP: mesmo padrão de colapso."])

doc.add_heading("Checagem de interação", level=2)
table(["Configuração · 5 folds", "Validação", "Gap", "Melhor época", "Parâmetros"],
      [rank_row(C["b2_final"][0], "Campeão do Bloco 2 · `4 blocos · same`")] +
      [rank_row(r, f"{'2º' if r['label'].startswith('rank2') else '3º'} do Bloco 1 · `3 blocos · {r['padding']}`") for r in C["chk_final"]],
      widths=[6, 3, 2.5, 2.5, 3], highlight=lambda i: i == 0)
bullets(["As duas ficaram abaixo do campeão de 4 blocos, e o Bloco 3 partiu dele. **Atenção:** a checagem foi feita sem regularização; o Complementar B mostrou que, com BatchNorm, a rede de 3 blocos passa à frente."])

doc.add_heading("Bloco 3 · BatchNorm, dropout e pooling", level=2)
question("Como a janela de pooling, o dropout nas camadas densas e a BatchNorm após cada convolução afetam a generalização? "
         "Grid: pooling {0, 2, 3} × dropout {0; 0,3; 0,5} × BatchNorm {não, sim} = 18 combinações, 6 válidas. 50 épocas, paciência 7.")
figure(FIG_HM_CNN_B3, "CNN, Bloco 3: acurácia de validação na triagem (pooling 2×2, único válido com 4 blocos).", width=Cm(10))
figure(FIG_CV["cnn_b3"], "CNN, Bloco 3: curvas por época com e sem BatchNorm (dropout 0,5, média dos folds 1–3).")
bullets(["**BatchNorm foi o maior ganho entre os hiperparâmetros da CNN: +3,3 a +4,2 p.p.**",
         "Sem BatchNorm, o dropout não faz diferença. Com BatchNorm, dropout 0,3 e 0,5 empatam no topo, e o 0,5 reduziu o gap de 0,142 para 0,095.",
         "O teste ficou ~1 p.p. abaixo da validação: dentro de 1 desvio, e esperado depois de escolher o melhor entre várias opções."])
facts([("Campeão", "BN · dropout 0,5 · pool 2"), ("Validação · 5 folds", "79,2% ± 1,3"), ("Teste", "78,3% ± 1,1"), ("Ganho sobre o Bloco 2", "+2,6 p.p.")])
note("**Ressalva.** A pergunta sobre a janela de pooling ficou sem resposta neste bloco: com 4 blocos, pooling 3×3 reduz o mapa a 0×0 e “sem pooling” gera 270 milhões de parâmetros. Por isso foi criado o Complementar B.")

doc.add_heading("Complementar B · Janela de pooling", level=2)
question("Blocos {2, 3, 4} × janela de pooling {2, 3, 4} com a receita campeã (kernel 3, same, BatchNorm, dropout 0,5, Adam 1e-3). A combinação 4 blocos + pooling 2×2 repete o campeão anterior.")
figure(FIG_HM_CNN_B, "CNN, Complementar B: acurácia de validação (5 folds).", width=Cm(10))
figure(FIG_CV["cnn_b"], "CNN, Complementar B: curvas por época (média dos 5 folds).")
pb = {p["exp"]: p for p in C["paired_B"]}
rows_b = []
for s in sorted(C["extraB_summary"], key=lambda s: -s["val"]):
    lab = s["label"]; b, p_ = int(lab[1]), int(lab[-1]); m = 32
    for _ in range(b):
        m //= p_
    pr = pb.get(s["exp"])
    rows_b.append([f"{b} blocos · pool {p_}×{p_}", f"{m}×{m}", params(s["params"]), pct(s["val"]) + " " + sd(s["val_std"]),
                   pct(s["test"]) + " " + sd(s["test_std"]), f"{pp(pr['delta'], 2)} · {pr['wins']}/5" if pr else "referência"])
table(["Configuração", "Mapa final", "Parâmetros", "Validação", "Teste", "vs 4 blocos · pool 2"], rows_b, widths=[3.6, 2, 2.2, 3, 3, 3.2], highlight=lambda i: i == 0, font=8.5)
bullets(["**Novo campeão da CNN:** 3 blocos com pooling 2×2 venceu a de 4 blocos em 4 de 5 folds (+1,0 p.p. na validação) e chegou a **79,5%** no teste.",
         "É a **interação entre blocos** na prática: sem regularização o 4º bloco ajudava; com BatchNorm, a rede de 3 blocos treina por mais tempo de forma produtiva (melhor época 12 contra 6).",
         "**O que importa é o tamanho final do mapa:** reduzir de menos (8×8) gera uma camada densa enorme e mais overfitting; reduzir demais (1×1) perde informação. O ponto ideal foi um mapa de 3–4 px alcançado gradualmente."])
callout("Reprodutibilidade", "A repetição do campeão anterior deu 78,7% de validação, contra 79,2% originais (−0,5 p.p., com as mesmas melhores épocas). "
        "Mesmo com seed fixa, operações da GPU (cuDNN) não são determinísticas: há um ruído de ~±0,5 p.p. entre execuções idênticas. "
        "Diferenças menores que ~1 p.p. só contam se repetirem fold a fold.")
slides(["`outputs_final/_relatorio/heatmap_cnn_b3b_pooling_conv_blocks_x_pool_size__val-accuracy_mean.png` e `…__num_parameters.png`",
        "Prova do novo campeão: `outputs_final/_relatorio/pareado_cnn_b3b_pooling__B4_pool2__val-accuracy.png` (pontos por fold)."])

# ---------------------------------------------------------------- comparação
doc.add_heading("MLP × CNN por classe", level=1)
figure(FIG_CLASSES, "Recall no teste por classe dos dois campeões sem augmentation (média de 5 folds), ordenado pelo ganho da CNN.", width=Cm(15))
table(["MLP · real", "previsto", "imagens", "CNN · real", "previsto", "imagens"],
      [[PT[m["real"]], PT[m["pred"]], str(m["n"]), PT[c["real"]], PT[c["pred"]], str(c["n"])] for m, c in zip(M["confusions_final"], C["confusions_final"])],
      widths=[3, 3, 2.2, 3, 3, 2.2], font=9)
para("Na MLP, confusões como **avião → navio** e **caminhão ↔ automóvel** indicam decisão guiada por **cor e fundo** (céu e mar azuis). "
     "Na CNN, esses erros saem do topo e o que sobra se concentra entre **animais parecidos** (cachorro ↔ gato, pássaro → cervo): a rede passou a olhar forma e textura.")

# ---------------------------------------------------------------- métricas por classe dos campeões
doc.add_heading("Métricas por classe e matrizes de confusão", level=1)
para("Precision (P), recall (R) e F1 de cada classe para os 13 melhores modelos: o campeão de cada etapa, as referências e os bônus. "
     "Treino e validação são a média dos 5 folds na melhor época (treino medido em modo avaliação, sem dropout e sem augmentation); "
     "o teste é a média dos 5 modelos, um por fold, nas 10.000 imagens de teste. As matrizes de confusão somam os 5 folds: "
     "50.000 previsões, 5.000 por classe real, com cada célula em % da linha.")
figure(FIG_F1_STAGES, "F1 no teste por classe em cada etapa (média de 5 folds). Contorno = campeã de cada rede sem augmentation.")
doc.add_heading("O que concluir", level=2)
bullets(D["class_conclusions"])
para("Nas tabelas a seguir, valores em %; a linha destacada é a classe de menor F1 no teste.", size=9, color="6f7889")
for net, title in (("mlp", "MLP"), ("cnn", "CNN")):
    doc.add_heading(f"Campeões da {title}", level=2)
    for ch in [c for c in CH if c["net"] == net]:
        doc.add_heading(f"{ch['label']} · {ch['sub']}", level=3)
        O = ch["overall"]
        para(f"`{ch['exp']}` · teste: acurácia **{pct(O['test']['accuracy']['mean'])} {sd(O['test']['accuracy']['std'])}**, "
             f"F1 macro {pct(O['test']['f1_macro']['mean'])}, loss {num(O['test']['loss']['mean'], 3)} · validação: {pct(O['val']['accuracy']['mean'])} · "
             f"melhor época {num(ch['best_epoch'])} · {params(ch['params'])} parâmetros", size=9, color="465063")
        class_table(ch)
        total = f"{sum(map(sum, ch['confusion'])):,}".replace(",", ".")
        figure(FIG_CM[ch["exp"]], f"Matriz de confusão no teste · {ch['label']} ({ch['sub']}). % de cada linha, soma dos {ch['cm_folds']} folds "
               f"({total} previsões). Diagonal = recall; azul = erros, mais escuro até 20% da linha.", width=Cm(11.5))
        bullets(ch["reading"], size=9.5)

# ---------------------------------------------------------------- bônus
doc.add_heading("Bônus · Data augmentation", level=1)
para("Pré-processamento, não hiperparâmetro da rede: por isso fica fora da sequência principal de blocos. **RandomCrop 32 com padding 4** "
     "(desloca a imagem até 4 pixels) e **flip horizontal** (p = 0,5), aplicados só no treino; validação e teste usam as imagens originais. "
     "Cada rede partiu do seu campeão, com a mesma receita **sem** augmentation re-treinada no mesmo bloco como referência pareada.")


def aug_table(net, name):
    A = AUG[net]
    pv = {p["exp"]: p for p in A["paired_val"]}
    rows = sorted(A["rows"], key=lambda r: -r["val"])
    table(["Configuração", "Folds", "Validação", "Teste", "Gap", "Melhor época", "vs sem augmentation"],
          [[name(r), str(r["folds"]), pct(r["val"]) + " " + sd(r["val_std"]), pct(r["test"]) + " " + sd(r["test_std"]), num(r["gap"], 3),
            num(r["best_epoch"]), "referência" if r["exp"] == A["ref"] else f"{pp(pv[r['exp']]['delta'], 2)} · {pv[r['exp']]['wins']}/{pv[r['exp']]['n']}"]
           for r in rows],
          widths=[4.6, 1.3, 2.6, 2.6, 1.5, 1.8, 2.6], highlight=lambda i: rows[i]["exp"] == A["champ"], font=8.5)


doc.add_heading("CNN · Augmentation × profundidade × dropout", level=2)
question("Quanto a diversidade extra de imagens acrescenta à CNN campeã? Com mais dados efetivos, a rede de 4 blocos volta a compensar? O dropout ainda é necessário? "
         "Grid: augmentation {não, sim} × blocos {3, 4} × dropout {0,3; 0,5} = 8 configurações. 100 épocas, paciência 10.")
figure(FIG_HM_AUG_CNN, "CNN, bônus: acurácia de validação sem e com augmentation (mesma escala de cor).", width=Cm(14))
aug_table("cnn", lambda r: f"{'com' if r['augment'] else 'sem'} aug · {r['conv_blocks']} blocos · dropout {r['dropout']:g}".replace(".", ","))
figure(FIG_CV["cnn_aug"], "CNN, bônus: curvas por época (média dos folds 1–3). Com augmentation, o treino segue útil por muito mais épocas.")
bullets(["**O maior salto de todo o estudo: +7,8 p.p. na validação** sobre a mesma receita sem augmentation, vencendo em **5 de 5 folds**; no teste, +7,4 p.p. (5 de 5).",
         "**A rede de 4 blocos voltou a vencer** (+0,8 p.p. sobre 3 blocos, em 3 de 3 folds). Sem augmentation, 3 e 4 blocos empatam. Com mais diversidade de dados, a capacidade extra deixa de virar overfitting: é a interação entre blocos no sentido inverso.",
         "O gap treino–validação caiu de 0,105 para 0,069, e a melhor época passou de ~10 para ~31: o modelo aprende por mais tempo sem decorar.",
         "O dropout **continuou útil**: 0,5 ≥ 0,3 com augmentation. A hipótese de que augmentation dispensaria o dropout não se confirmou."])
facts([("Campeão com augmentation", "4 blocos · dropout 0,5"), ("Validação · 5 folds", "87,2% ± 0,5"), ("Teste", "86,4% ± 0,6"), ("Ganho pareado (teste)", "+7,4 p.p. · 5/5")])
note("**Ressalvas.** A receita sem augmentation, re-treinada com 100 épocas e paciência 10, deu 79,4% de validação contra 79,8% no Complementar B: "
     "−0,4 p.p., dentro do ruído entre execuções. Nenhuma configuração com augmentation chegou perto do limite de 100 épocas (melhores épocas entre 18 e 53).")
slides(["`outputs_augmentation/_relatorio/heatmap_cnn_b4_augmentation_conv_blocks_x_cnn_dropout__val-accuracy_mean.png`",
        "Prova do ganho: `outputs_augmentation/_relatorio/pareado_cnn_b4_augmentation__aug0_B3_do0.5__val-accuracy.png` e as curvas por época."])

doc.add_heading("MLP · Augmentation × dropout", level=2)
question("Deslocamentos e espelhamentos ajudam uma rede que não tem noção de vizinhança entre pixels? Grid: augmentation {não, sim} × dropout {0; 0,2} = 4 configurações, 5 folds. 150 épocas, paciência 15.")
figure(FIG_HM_AUG_MLP, "MLP, bônus: acurácia de validação (5 folds).", width=Cm(10))
aug_table("mlp", lambda r: f"{'com' if r['augment'] else 'sem'} aug · dropout {r['dropout']:g}".replace(".", ","))
figure(FIG_CV["mlp_aug"], "MLP, bônus: curvas por época (média dos 5 folds). Com augmentation, a MLP precisa de ~140 épocas.")
bullets(["**Ganho real, mas pequeno: +2,9 p.p. na validação** (5 de 5 folds) e +2,7 p.p. no teste, contra +7,4 p.p. da CNN.",
         "**Convergência ~9× mais lenta:** a melhor época passou de ~15 para ~141. Para a MLP, cada imagem deslocada é uma entrada nova, e ela precisa “ver” cada posição separadamente.",
         "O gap caiu de 0,148 para 0,052, e dropout 0,2 continuou ajudando (+2,0 p.p. sobre augmentation sem dropout).",
         "Por classe, o ganho é irregular: caminhão, cavalo e gato melhoram ~8 p.p., mas pássaro e cervo pioram ~3 p.p."])
facts([("Campeão com augmentation", "4 × 256 · dropout 0,2"), ("Validação · 5 folds", "57,9% ± 0,3"), ("Teste", "57,8% ± 0,2"), ("Ganho pareado (teste)", "+2,7 p.p. · 5/5")])
callout("Ressalva importante", "O campeão da MLP com augmentation foi **limitado pelo orçamento**: a melhor época ficou entre 127 e 148 num limite de 150, ainda subindo. "
        "57,8% é um **piso**; com mais épocas o ganho provavelmente seria um pouco maior. A receita sem augmentation reproduziu exatamente o resultado anterior (55,0% de validação).")
slides(["`outputs_augmentation/_relatorio/heatmap_mlp_b4_augmentation_dropout_x_augment__melhor_epoca_media.png` mostra o custo em épocas.",
        "Mensagem: “o mesmo pré-processamento vale +7,4 p.p. na CNN e +2,7 p.p. na MLP — a diferença é a estrutura espacial”."])

# ---------------------------------------------------------------- hiperparâmetros
doc.add_heading("Conclusões por hiperparâmetro", level=1)
para("Os blocos respondem “qual configuração venceu”; esta seção responde “o que cada hiperparâmetro faz e por quê”. "
     "Para cada um: o efeito medido na validação, o resultado por valor testado, a evidência e o mecanismo que explica o resultado.")
table(["Hiperparâmetro", "Rede", "Escolha", "Efeito na validação", "Veredito"],
      [[f"**{h['name']}**", " · ".join(h["nets"]), h["choice"], f"**{h['effect']}** · {h['effect_note']}", h["verdict"]] for h in HY],
      widths=[3.6, 1.6, 4.0, 5.0, 2.8], font=8.5)
for group in dict.fromkeys(h["group"] for h in HY):
    doc.add_heading(group, level=2)
    for h in [x for x in HY if x["group"] == group]:
        doc.add_heading(h["name"], level=3)
        facts([("Rede", " · ".join(h["nets"])), ("Valores", h["values"]), ("Escolha", h["choice"]), ("Efeito", f"{h['effect']} ({h['effect_note']})"), ("Veredito", h["verdict"])])
        if h["marg"]:
            table(["Valor", "Validação", "Melhor época", "Detalhe"],
                  [[("**" + m["label"] + "** (escolha)") if m["champ"] else m["label"], pct(m["mean"]), num(m["ep"]),
                    f"média de {m['n']} configurações · melhor {pct(m['best'])}" if m["n"] > 1 else "—"] for m in h["marg"]],
                  widths=[5.2, 2.6, 2.6, 6.6], font=8.5, highlight=lambda i, h=h: h["marg"][i]["champ"])
            para(h["marg_note"], size=8.5, color="6f7889")
        bullets(h["evidence"], size=9.5)
        callout("Por quê", h["why"])

# ---------------------------------------------------------------- lições
doc.add_heading("Lições metodológicas", level=1)
table(["Lição", "Evidência", "Como dizer no slide"], [
    ["**Escolher pela validação, nunca pelo teste**", "Validação e teste dos campeões finais concordam: MLP 55,0% × 55,2%; CNN 79,8% × 79,5%.", "“O teste foi usado uma única vez, só nos campeões.”"],
    ["**Maldição do vencedor**", "1º da triagem do Bloco 1 da MLP: 53,1% em 3 folds → 52,7% em 5.", "“O melhor entre muitos é otimista; por isso as finalistas completam os 5 folds.”"],
    ["**Interação entre blocos**", "CNN: 4 blocos venciam sem regularização; com BatchNorm, 3 blocos vence (4/5 folds); com augmentation, 4 blocos volta a vencer (3/3).", "“A melhor profundidade depende da regularização e dos dados.”"],
    ["**Orçamento de épocas**", "SGD com lr baixo e dropout 0,5 esbarraram no limite; com 100 épocas, o dropout 0,5 continuou pior (0/5). A MLP com augmentation também esbarrou (época ~141 de 150).", "“Verificamos se os perdedores só precisavam de mais tempo.”"],
    ["**Ruído entre execuções**", "Mesma configuração repetida: −0,4 a −0,5 p.p. de diferença na CNN.", "“Diferenças abaixo de ~1 p.p. só valem se repetirem fold a fold.”"],
    ["**Dados × arquitetura**", "Augmentation rendeu +7,4 p.p. à CNN (maior salto do estudo) e +2,7 p.p. à MLP.", "“O mesmo pré-processamento só é bem aproveitado por quem enxerga estrutura espacial.”"],
    ["**Comparação pareada**", "Entropia cruzada × MSE: +0,34 p.p., mas em 5 de 5 folds.", "“Mesmas partições permitem comparar fold a fold.”"],
], widths=[4.2, 7, 5.8], font=9)

# ---------------------------------------------------------------- roteiro
doc.add_heading("Roteiro de slides", level=1)
outline = [
    ("Título e equipe", "Estudo de ablação de MLP e CNN no CIFAR-10.", "—"),
    ("Problema e dados", "CIFAR-10: 60 mil imagens 32×32, 10 classes; objetivo de justificar cada escolha de arquitetura.", "Grade de imagens de exemplo (a produzir)"),
    ("Metodologia", "5 folds com partições fixas, grid search em blocos, triagem + confirmação, escolha pela validação, teste revelado no fim.", "Diagrama do fluxo dos blocos (a produzir)"),
    ("Infraestrutura e registro", "Kaggle 2× T4, dados na GPU, W&B + CSV por época, MD5 do dataset, retomada.", "Print do projeto no W&B"),
    ("Referências (B0)", "Linear 38,8%, MLP original 51,1%, LeNet 69,7% no teste.", "Tabelas dos Blocos 0"),
    ("MLP · Topologia (B1)", "Profundidade ajuda até ~2–3 camadas; largura acima de 256 não; empate estatístico no topo.", "Figuras do Bloco 1 da MLP"),
    ("MLP · Otimização (B2)", "Faixas úteis de lr por algoritmo; colapso do Adam com lr alto; Adam 3e-4 vence.", "Figuras do Bloco 2 da MLP"),
    ("MLP · Regularização e erro (B3 + A)", "Dropout 0,2 = +1,5 p.p.; dropout atrasa o overfitting; 0,5 subajusta (5/5 folds); entropia cruzada > MSE (5/5).", "Heatmap + curvas do Bloco 3"),
    ("MLP · Resultado", "55,2% no teste (+4,0 p.p.); teto da arquitetura; confusões por cor e fundo.", "Evolução da MLP"),
    ("CNN · Topologia (B1)", "Max pooling vence stride 2 em todas as combinações; 4 blocos, kernel 3.", "Figuras do Bloco 1 da CNN"),
    ("CNN · Otimização (B2)", "Sem ganho: Adam padrão já adequado; mesmo padrão de colapso da MLP.", "Heatmap do Bloco 2 da CNN"),
    ("CNN · BatchNorm e dropout (B3)", "BatchNorm +3,3 a +4,2 p.p.; dropout só ajuda com BatchNorm; pergunta do pooling em aberto.", "Heatmap + curvas do Bloco 3"),
    ("CNN · Pooling (+B)", "Novo campeão 3 blocos + pool 2 (79,5%); tamanho final do mapa; reprodutibilidade.", "Heatmap, curvas e comparação pareada"),
    ("MLP × CNN", "+24,3 p.p. no teste; CNN vence todas as classes; muda o tipo de erro.", "Recall por classe"),
    ("Bônus · Augmentation", "Mesmo pré-processamento: CNN +7,4 p.p. (86,4%, 4 blocos voltam a vencer) × MLP +2,7 p.p. (57,8%, limitada por épocas).", "Heatmaps e curvas do bônus"),
    ("Métricas por classe", "Gato é a pior classe e automóvel/navio as melhores nos 13 campeões; gato ↔ cachorro domina as confusões; a CNN ganha mais nos animais.", "F1 por etapa + matriz de confusão do campeão"),
    ("Hiperparâmetros: o que pesou", "Decisivos: augmentation (CNN), max pooling × stride, BatchNorm. Moderados: profundidade, kernel, janela de pooling, dropout. Empate: otimizador. Maior risco: learning rate.", "Tabela-resumo de hiperparâmetros"),
    ("Lições metodológicas", "Validação × teste, maldição do vencedor, interação, orçamento, ruído, dados × arquitetura, pareamento.", "Tabela de lições"),
    ("Conclusões", "O que mais pesou em cada rede e por quê; limitações e próximos passos.", "—"),
]
table(["#", "Slide", "Conteúdo", "Figura"], [[f"{i:02d}", f"**{a}**", b, c] for i, (a, b, c) in enumerate(outline, 1)], widths=[1, 4, 7.5, 4.5], font=8.5)

# ---------------------------------------------------------------- arquivos
doc.add_heading("Onde está cada coisa", level=1)
table(["Recurso", "Local"], [
    ["W&B", "wandb.ai/models-universidade-federal-de-pernambuco/cifar10-rn · uma run por fold (`{experimento}-foldK`), agrupadas pelo nome do experimento. Para ver médias entre folds, agrupe por Group."],
    ["Resultados locais", "`outputs_mlp/`, `outputs_cnn/`, `outputs_final/` (complementares) e `outputs_augmentation/` (bônus) no repositório `~/rn`."],
    ["Figuras prontas", "`outputs_*/_relatorio/*.png`. As figuras deste documento (inclusive as matrizes de confusão somadas por fold) estão em `relatorio/figuras/`."],
    ["Rankings e campeões", "`outputs_*/_grids/{bloco}/`: `ranking_triagem.csv`, `ranking_final.csv`, `campeao.json`."],
    ["Notebooks", "`Kaggle_MLP.ipynb` e `Kaggle_CNN.ipynb`: cada um executa o estudo completo da rede (blocos, complementar e bônus) em uma única execução."],
    ["Código", "github.com/diegoflyra/dfal-neural-networks (privado): `src/`, `grids/`, `tests/`, `tools/relatorio/` (gera este documento e o caderno)."],
], widths=[3.6, 13.4], font=9)
para("Percentuais com vírgula; “p.p.” = pontos percentuais; “±” = desvio padrão entre folds. Números extraídos dos arquivos de resultados, idênticos às métricas registradas no W&B.", size=8.5, color="6f7889")

out = "miniprojeto 1 doc.docx"
doc.save(out)
print("docx:", out, os.path.getsize(out) // 1024, "KB | figuras:", fig_counter[0])
